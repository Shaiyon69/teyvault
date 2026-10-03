"""Wish history: locate the authkey URL, fetch the log, compute pity and 50/50 stats, UIGF I/O."""
import calendar
import json
import os
import re
import time
import urllib.parse
from importlib import resources
from pathlib import Path

from teyvat import db
from teyvat.hoyolab import ApiError, request

API_HOSTS = {
    "hk4e_global": "https://public-operation-hk4e-sg.hoyoverse.com/gacha_info/api/getGachaLog",
    "hk4e_cn": "https://public-operation-hk4e.mihoyo.com/gacha_info/api/getGachaLog",
}
# 400 (second character event banner) is returned by the 301 query, so it is not listed.
QUERY_TYPES = ("100", "200", "301", "302", "500")
POOLS = {
    "character": (("301", "400"), 90),
    "weapon": (("302",), 80),
    "standard": (("200",), 90),
    "chronicled": (("500",), 90),
    "beginner": (("100",), 90),
}
RETCODE_AUTHKEY_TIMEOUT = -101
RETCODE_TOO_FREQUENT = -110
PAGE_DELAY_S = 1  # gentle on the API; the in-game page loads one page per click
# paimon.moe's open-source image set; file name = lowercased name, quotes dropped, spaces -> "_".
ICON_CDN = "https://cdn.jsdelivr.net/gh/MadeBaruna/paimon-moe@main/static/images"

URL_RE = re.compile(rb"https://[^\x00\s\"']+?authkey=[^\x00\s\"']+")
GAME_DIR_RE = re.compile(r"([A-Za-z]:/[^\n\r:]*?(?:GenshinImpact_Data|YuanShen_Data))")


# --- Authkey location -------------------------------------------------------

def log_paths(base=None) -> list[Path]:
    """Candidate game log files. `base` overrides %USERPROFILE% (e.g. a Wine prefix user dir)."""
    root = Path(base or os.environ.get("USERPROFILE", Path.home())) / "AppData/LocalLow/miHoYo"
    return [root / game / name for game in ("Genshin Impact", "原神")
            for name in ("output_log.txt", "Player.log")]


def find_game_data_dir(log_text: str) -> Path | None:
    matches = GAME_DIR_RE.findall(log_text)
    return Path(matches[-1]) if matches else None


def extract_gacha_url(cache_bytes: bytes) -> str | None:
    """Return the most recent wish-history URL in a webCaches data_2 blob."""
    urls = [m.decode("utf-8", "ignore") for m in URL_RE.findall(cache_bytes)]
    urls = [u for u in urls if "gacha" in u]
    return urls[-1] if urls else None


def locate_gacha_url(base=None) -> str:
    for log in log_paths(base):
        if not log.exists():
            continue
        data_dir = find_game_data_dir(log.read_text("utf-8", "ignore"))
        if not data_dir:
            continue
        caches = sorted((data_dir / "webCaches").glob("*/Cache/Cache_Data/data_2"),
                        key=lambda p: p.stat().st_mtime)
        for cache in reversed(caches):
            try:
                url = extract_gacha_url(cache.read_bytes())
            except PermissionError:
                raise SystemExit(f"Cannot read {cache} (locked). Close the game and retry.")
            if url:
                return url
    raise SystemExit("No wish history URL found. Open the wish history page in game, then retry "
                     "(or pass --url).")


# --- Fetching ---------------------------------------------------------------

def api_params(gacha_url: str) -> tuple[str, dict]:
    """Turn a cached page URL into (API endpoint, base query params)."""
    query = urllib.parse.urlsplit(gacha_url.split("#")[0]).query
    params = {k: v[-1] for k, v in urllib.parse.parse_qs(query).items()}
    if "authkey" not in params:
        raise SystemExit("URL has no authkey.")
    biz = params.get("game_biz", "hk4e_global")
    endpoint = API_HOSTS.get(biz)
    if not endpoint:
        raise SystemExit(f"Unsupported game_biz {biz!r}.")
    params.update(authkey_ver="1", lang="en")
    return endpoint, params


def to_row(item: dict) -> dict:
    return {"id": item["id"], "uid": item["uid"], "gacha_type": item["gacha_type"],
            "item_type": item.get("item_type"), "name": item["name"],
            "rank": int(item["rank_type"]), "time": item["time"]}


def fetch_type(endpoint, params, gacha_type, known, *, sleep=time.sleep, log=print) -> list[dict]:
    """Fetch one banner type newest-first, stopping at the first id `known(id)` reports."""
    rows, end_id, page, retries = [], "0", 1, 0
    while True:
        try:
            payload = request(endpoint, params={**params, "gacha_type": gacha_type,
                                                "page": page, "size": 20, "end_id": end_id})
        except ApiError as e:
            if e.retcode == RETCODE_TOO_FREQUENT and retries < 5:
                retries += 1
                sleep(2)
                continue
            if e.retcode == RETCODE_AUTHKEY_TIMEOUT:
                raise SystemExit("Authkey expired. Open wish history in game again, then retry.")
            raise
        items = payload["data"]["list"]
        for item in items:
            if known(item["id"]):
                log(f"  type {gacha_type}: {len(rows)} new ({page} pages)")
                return rows
            rows.append(to_row(item))
        if len(items) < 20:
            log(f"  type {gacha_type}: {len(rows)} new ({page} pages)")
            return rows
        end_id, page = items[-1]["id"], page + 1
        sleep(PAGE_DELAY_S)


def sync(conn, gacha_url, **kw) -> int:
    endpoint, params = api_params(gacha_url)
    added = 0
    for gacha_type in QUERY_TYPES:
        rows = fetch_type(endpoint, params, gacha_type, lambda i: db.has_wish(conn, i), **kw)
        added += db.insert_wishes(conn, rows)
    return max(0, added - db.drop_covered_synthetic(conn))


# --- Stats ------------------------------------------------------------------

def load_standard() -> dict[str, str]:
    text = resources.files("teyvat").joinpath("data/standard_5stars.json").read_text("utf-8")
    return json.loads(text)["characters"]


def is_standard(name, pull_time, standard) -> bool:
    since = standard.get(name)
    return since is not None and pull_time[:10] >= since


def pool_stats(rows, hard_pity, standard=None) -> dict:
    """rows: one pity pool in pull order. `standard` given => track 50/50 (character pool)."""
    pity, five_stars, four_stars, guaranteed = 0, [], [], False
    won = lost = pity4 = 0
    for r in rows:
        pity += 1
        pity4 += 1
        if r["rank"] == 4:
            four_stars.append({"id": r["id"], "rank": 4, "name": r["name"], "item_type": r["item_type"],
                               "pity": pity4, "time": r["time"], "outcome": None})
            pity4 = 0
        if r["rank"] != 5:
            continue
        outcome = None
        if standard is not None:
            off_banner = is_standard(r["name"], r["time"], standard)
            if guaranteed:
                outcome, guaranteed = "guaranteed", False
            elif off_banner:
                outcome, guaranteed, lost = "lost", True, lost + 1
            else:
                outcome, won = "won", won + 1
        five_stars.append({"id": r["id"], "rank": 5, "name": r["name"], "item_type": r["item_type"],
                           "pity": pity, "time": r["time"], "outcome": outcome})
        pity = 0
    return {"total": len(rows), "pity": pity, "hard_pity": hard_pity, "five_stars": five_stars,
            "won": won, "lost": lost, "guaranteed": guaranteed, "pity4": pity4,
            "four_stars": four_stars}


def stats(rows, standard) -> dict:
    out = {}
    for pool, (types, hard) in POOLS.items():
        pool_rows = [r for r in rows if r["gacha_type"] in types]
        if pool_rows:
            out[pool] = pool_stats(pool_rows, hard, standard if pool == "character" else None)
    return out


def slug(name) -> str:
    """paimon.moe's id for a character/weapon name ("Kuki Shinobu" -> "kuki_shinobu")."""
    return re.sub(r"[^a-z0-9-]+", "_", re.sub(r"['\"]", "", name.lower())).strip("_")


def icon_url(name, item_type) -> str:
    folder = "weapons" if item_type == "Weapon" else "characters"
    return f"{ICON_CDN}/{folder}/{slug(name)}.png"


# --- UIGF import/export -----------------------------------------------------

def uigf_timezone(uid: str) -> int:
    return {"6": -5, "7": 1}.get(uid[0], 8)


def export_uigf(conn) -> dict:
    accounts = []
    for uid in db.uids(conn):
        accounts.append({"uid": uid, "timezone": uigf_timezone(uid), "lang": "en-us", "list": [
            {"uigf_gacha_type": "301" if r["gacha_type"] == "400" else r["gacha_type"],
             "gacha_type": r["gacha_type"], "item_id": "", "count": "1", "time": r["time"],
             "name": r["name"], "item_type": r["item_type"], "rank_type": str(r["rank"]),
             "id": r["id"]}
            for r in db.wishes(conn, uid)]})
    return {"info": {"export_timestamp": int(time.time()), "export_app": "Teyvault",
                     "export_app_version": "0.1.0", "version": "v4.0"}, "hk4e": accounts}


def import_uigf(conn, data: dict) -> int:
    """Accepts UIGF v4 ({"hk4e": [...]}) and v2/v3 ({"info": {"uid"}, "list": [...]})."""
    if "hk4e" in data:
        accounts = data["hk4e"]
    else:
        accounts = [{"uid": data["info"]["uid"], "list": data["list"]}]
    added = 0
    for acc in accounts:
        rows = [to_row({**item, "uid": str(item.get("uid") or acc["uid"])}) for item in acc["list"]]
        added += db.insert_wishes(conn, rows)
    return added


# --- Excel import/export ----------------------------------------------------

XLSX_COLUMNS = ("Time", "Name", "Type", "Rarity", "Pity", "Banner", "UID", "ID")
XLSX_FILLS = {5: "FFE6C07B", 4: "FFD9C7F7"}


def export_xlsx(conn) -> bytes:
    """One sheet per banner like paimon.moe's export, 5★/4★ rows highlighted. Keeps ids, so the
    file imports back without loss. Ids and UIDs are text: Excel would round 19-digit numbers."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for pool, (types, _) in POOLS.items():
        ws = None
        for uid in db.uids(conn):
            pity = 0
            for r in db.wishes(conn, uid):
                if r["gacha_type"] not in types:
                    continue
                if ws is None:
                    ws = wb.create_sheet(pool.title())
                    ws.append(XLSX_COLUMNS)
                    for cell in ws[1]:
                        cell.font = Font(bold=True)
                    ws.freeze_panes = "A2"
                    for col, width in zip("ABCDEFGH", (20, 32, 11, 8, 6, 8, 12, 22)):
                        ws.column_dimensions[col].width = width
                pity += 1
                ws.append([r["time"], r["name"], r["item_type"], r["rank"], pity, r["gacha_type"],
                           uid, r["id"]])
                if fill := XLSX_FILLS.get(r["rank"]):
                    for cell in ws[ws.max_row]:
                        cell.fill = PatternFill("solid", fgColor=fill)
                if r["rank"] == 5:
                    pity = 0
    if not wb.sheetnames:
        wb.create_sheet("Wishes").append(XLSX_COLUMNS)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# paimon.moe sheet name (lowercased, first word) -> gacha_type. Its export has no ids or UID.
PAIMON_SHEETS = {"character": "301", "weapon": "302", "standard": "200", "beginners'": "100",
                 "chronicled": "500"}


def paimon_id(time_str, gacha_type, roll) -> str:
    """Synthetic id for a pull without one: 18 digits (real ids have 19, so these always sort
    before real ones), ordered by time then paimon.moe's per-banner roll number. Stable, so
    re-importing the same file adds nothing."""
    ts = calendar.timegm(time.strptime(time_str, "%Y-%m-%d %H:%M:%S"))
    return str(ts * 10**8 + QUERY_TYPES.index(gacha_type) * 10**6 + roll)


def import_xlsx(conn, data: bytes, uid=None) -> int:
    """Reads Teyvault exports (header row with XLSX_COLUMNS, any order) and paimon.moe exports.
    paimon.moe files carry no UID: `uid` names the account, else the only UID in the DB."""
    from io import BytesIO
    from openpyxl import load_workbook

    def text_time(t):  # Excel may hand back a datetime if the cell was typed
        return t if isinstance(t, str) else t.strftime("%Y-%m-%d %H:%M:%S")

    rows, paimon = [], False
    for ws in load_workbook(BytesIO(data), read_only=True).worksheets:
        values = ws.iter_rows(values_only=True)
        header = next(values, None) or ()
        col = {h: i for i, h in enumerate(header)}
        if {"ID", "UID", "Banner", "Name", "Rarity", "Time"} <= set(header):
            for v in values:
                if v[col["ID"]] is None:
                    continue
                rows.append({"id": str(v[col["ID"]]), "uid": str(v[col["UID"]]),
                             "gacha_type": str(v[col["Banner"]]), "name": v[col["Name"]],
                             "item_type": v[col["Type"]] if "Type" in col else None,
                             "rank": int(v[col["Rarity"]]), "time": text_time(v[col["Time"]])})
        elif {"Type", "Name", "Time", "⭐", "#Roll"} <= set(header):
            gacha_type = PAIMON_SHEETS.get(ws.title.lower().split()[0])
            if gacha_type is None:
                continue
            if uid is None:
                known = db.uids(conn)
                if len(known) != 1:
                    raise SystemExit("This paimon.moe file has no UID. Sign in to HoYoLAB first so "
                                     "Teyvault knows which account it belongs to.")
                uid = known[0]
            paimon = True
            for v in values:
                if v[col["Name"]] is None:
                    continue
                t = text_time(v[col["Time"]])
                rows.append({"id": paimon_id(t, gacha_type, int(v[col["#Roll"]])), "uid": str(uid),
                             "gacha_type": gacha_type, "name": v[col["Name"]],
                             "item_type": v[col["Type"]], "rank": int(v[col["⭐"]]), "time": t})
    if not rows:
        raise SystemExit("No wish sheets found. Use a Teyvault or paimon.moe Excel export.")
    added = db.insert_wishes(conn, rows)
    return max(0, added - db.drop_covered_synthetic(conn)) if paimon else added
